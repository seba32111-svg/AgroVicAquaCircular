"""Genera diagramas, informe, presentación, guion y paquete final."""
from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import textwrap
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor as DocxColor
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches as PInches, Pt as PPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image as PdfImage,
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
EVIDENCE = ARTIFACTS / "evidence"
DIAGRAMS = ARTIFACTS / "diagramas"
DIST = ROOT / "dist"
SUMMARY_PATH = ARTIFACTS / "test_results.json"

FOREST = (15, 66, 58)
FOREST_DARK = (9, 43, 41)
AQUA = (25, 184, 174)
LIME = (167, 214, 90)
PAPER = (248, 250, 246)
INK = (16, 39, 43)
MUTED = (94, 116, 112)

RF = [
    ("RF01", "Inicio de sesión y recuperación de acceso mediante solicitud temporal."),
    ("RF02", "Autorización diferenciada por rol, ruta y acción."),
    ("RF03", "CRUD de usuarios y operadores: crear, consultar, modificar y desactivar/reactivar registros."),
    ("RF04", "CRUD de zonas de cultivo: alta, consulta, edición y baja lógica, preservando el historial."),
    ("RF05", "CRUD de dispositivos ESP32, provisión individual de token y administración de fuentes y sensores."),
    ("RF06", "Recepción de humedad, lluvia, niveles, caudal y calidad cada cinco minutos o menos."),
    ("RF07", "Registro del volumen capturado o recuperado."),
    ("RF08", "Prioridad de agua recuperada solo con nivel y calidad aptos."),
    ("RF09", "Bloqueo o detención segura por nivel, calidad, sensor, timeout o desconexión."),
    ("RF10", "Riego manual autorizado con fuente, motivo, volumen y duración."),
    ("RF11", "Umbrales, horario, duración, fuente y calidad configurables."),
    ("RF12", "Panel con métricas, estado de dispositivos y alertas."),
    ("RF13", "Alertas por humedad, saturación, calidad, nivel, desconexión y falla."),
    ("RF14", "Historial y exportación por zona, fuente y periodo."),
    ("RF15", "Auditoría de accesos, cambios, decisiones y acciones."),
]
RNF = [
    ("RNF01", "Contraseñas protegidas con bcrypt."),
    ("RNF02", "HTTPS en despliegue y token de dispositivo almacenado como hash."),
    ("RNF03", "Autorización de servidor para cada recurso protegido."),
    ("RNF04", "Validación de datos, rangos e idempotencia."),
    ("RNF05", "Tiempo p95 del panel inferior a tres segundos en el piloto."),
    ("RNF06", "Meta de disponibilidad de 99% y detención local ante falla."),
    ("RNF07", "Interfaz utilizable desde 360 px."),
    ("RNF08", "Compatibilidad con navegadores vigentes."),
    ("RNF09", "Modelo escalable por zona y dispositivo."),
    ("RNF10", "Exportación, respaldo, modularidad y documentación."),
]
OBJECTIVES = [
    ("OE1", "Problema, actores, 15 RF y 10 RNF documentados.", "Cumplido"),
    ("OE2", "Arquitectura, diagramas y modelo NoSQL de al menos ocho colecciones.", "Cumplido: 10 colecciones"),
    ("OE3", "Integración ESP32 con lecturas cada cinco minutos o menos.", "Cumplido en firmware y simulador; terreno pendiente"),
    ("OE4", "Web responsive, dos perfiles, tres CRUD completos, panel y control.", "Cumplido"),
    ("OE5", "Controles de seguridad, recuperación y auditoría.", "Cumplido en prototipo"),
    ("OE6", "500 lecturas, 10 ciclos de reúso e indicadores.", "Cumplido: 500, 10 y 67,01%"),
    ("OE7", "16 casos, al menos 90% aprobados y versión reproducible.", "Determinado por la ejecución adjunta"),
]
COLLECTIONS = [
    ("usuarios", "Identidad, rol, estado y hash"),
    ("zonas_cultivo", "Cultivo, ubicación y superficie"),
    ("dispositivos", "ESP32, estado y hash de token"),
    ("lecturas_sensor", "Serie temporal de telemetría"),
    ("fuentes_agua", "Principal y recuperada por zona"),
    ("reglas_riego", "Umbrales y versiones"),
    ("eventos_riego", "Ciclos, fuente y volumen"),
    ("alertas", "Severidad y flujo de atención"),
    ("auditoria", "Acceso, cambios y decisiones"),
    ("tokens_recuperacion", "Vigencia y uso único"),
]
API = [
    ("POST", "/api/v1/auth/login", "Sesión por perfil"),
    ("POST", "/api/v1/auth/recuperaciones", "Recuperación genérica"),
    ("POST", "/api/v1/lecturas", "Telemetría autenticada"),
    ("GET", "/api/v1/zonas/{id}/estado", "Estado de zona"),
    ("GET/POST", "/api/v1/riegos", "Consulta y orden manual"),
    ("PATCH", "/api/v1/alertas/{id}/confirmacion", "Atención de alerta"),
    ("GET", "/api/v1/reportes/balance", "Indicadores hídricos"),
    ("POST", "/api/v1/reglas", "Regla versionada"),
    ("POST/PATCH", "/api/v1/dispositivos", "Provisionamiento y estado"),
]


def load_summary():
    if not SUMMARY_PATH.exists():
        raise SystemExit("No existe artifacts/test_results.json. Ejecuta pytest primero.")
    return json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))


def font(size=28, bold=False):
    choices = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for path in choices:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def wrapped(draw, text, box, *, size=26, color=INK, bold=False, spacing=8, align="left"):
    x, y, width, _height = box
    chars = max(12, int(width / (size * 0.58)))
    lines = textwrap.wrap(text, chars)
    draw.multiline_text(
        (x, y),
        "\n".join(lines),
        font=font(size, bold),
        fill=color,
        spacing=spacing,
        align=align,
    )


def arrow(draw, start, end, color=AQUA, width=6):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 18, y - 10), (x - 18, y + 10)], fill=color)


def architecture_diagram():
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 1050), PAPER)
    draw = ImageDraw.Draw(image)
    draw.text((90, 60), "Arquitectura AgroVic AquaCircular", font=font(50, True), fill=INK)
    draw.text((90, 125), "Capas, decisiones de seguridad y trazabilidad", font=font(24), fill=MUTED)
    boxes = [
        ((90, 260, 420, 630), "CAMPO / BORDE", "ESP32\nSensores de humedad, nivel,\ncaudal y calidad\n\nWatchdog local\nVálvulas y bomba", FOREST_DARK),
        ((535, 260, 865, 630), "API Y REGLAS", "Flask / JSON\nValidación e idempotencia\n\nSelección de fuente\nControl de permisos\nAuditoría", FOREST),
        ((980, 260, 1310, 630), "DATOS", "MongoDB 7\n10 colecciones\n\n500 lecturas\n14 ciclos de riego\nExportación JSON", (21, 110, 99)),
        ((1425, 260, 1755, 630), "EXPERIENCIA", "Panel responsive\nAdministrador / Operador\n\nAlertas\nBalance hídrico\nReportes CSV", (28, 134, 118)),
    ]
    for rectangle, title, body, fill in boxes:
        draw.rounded_rectangle(rectangle, radius=26, fill=fill)
        x1, y1, x2, y2 = rectangle
        draw.text((x1 + 28, y1 + 28), title, font=font(20, True), fill=LIME)
        wrapped(draw, body, (x1 + 28, y1 + 85, x2 - x1 - 56, y2 - y1 - 110), size=24, color=(235, 247, 242), spacing=13)
    arrow(draw, (420, 445), (535, 445))
    arrow(draw, (865, 445), (980, 445))
    arrow(draw, (1310, 445), (1425, 445))
    draw.rounded_rectangle((230, 755, 1570, 930), radius=24, fill=(229, 246, 239), outline=(181, 222, 210), width=3)
    draw.text((280, 790), "Principio transversal", font=font(19, True), fill=FOREST)
    wrapped(
        draw,
        "Si la calidad es inválida, el nivel es insuficiente, el sensor falla o expira el tiempo máximo, el sistema bloquea el reúso y conduce a un estado seguro.",
        (280, 835, 1240, 90),
        size=25,
        color=INK,
    )
    path = DIAGRAMS / "Arquitectura_AgroVic_AquaCircular.png"
    image.save(path, quality=95)
    return path


def data_model_diagram():
    image = Image.new("RGB", (1800, 1180), (250, 251, 248))
    draw = ImageDraw.Draw(image)
    draw.text((80, 55), "Modelo NoSQL y relaciones lógicas", font=font(48, True), fill=INK)
    positions = {
        "usuarios": (90, 230), "zonas_cultivo": (630, 185), "dispositivos": (1170, 230),
        "auditoria": (90, 555), "reglas_riego": (630, 510), "lecturas_sensor": (1170, 555),
        "tokens_recuperacion": (90, 880), "fuentes_agua": (630, 835), "eventos_riego": (1170, 835),
        "alertas": (1450, 555),
    }
    relations = [
        ("usuarios", "auditoria"), ("usuarios", "tokens_recuperacion"),
        ("zonas_cultivo", "dispositivos"), ("zonas_cultivo", "reglas_riego"),
        ("zonas_cultivo", "fuentes_agua"), ("zonas_cultivo", "eventos_riego"),
        ("dispositivos", "lecturas_sensor"), ("lecturas_sensor", "alertas"),
        ("eventos_riego", "auditoria"), ("alertas", "auditoria"),
    ]
    centers = {}
    for name, (x, y) in positions.items():
        w = 280 if name != "lecturas_sensor" else 300
        h = 150
        centers[name] = (x + w // 2, y + h // 2)
    for source, target in relations:
        draw.line([centers[source], centers[target]], fill=(177, 205, 197), width=4)
    details = dict(COLLECTIONS)
    for name, (x, y) in positions.items():
        w = 280 if name != "lecturas_sensor" else 300
        draw.rounded_rectangle((x, y, x + w, y + 150), radius=18, fill=(255, 255, 255), outline=(156, 202, 191), width=3)
        draw.text((x + 18, y + 20), name, font=font(21, True), fill=FOREST)
        wrapped(draw, details[name], (x + 18, y + 62, w - 36, 70), size=17, color=MUTED, spacing=5)
    draw.rounded_rectangle((1375, 890, 1730, 1080), radius=20, fill=FOREST_DARK)
    draw.text((1400, 920), "Volumen del piloto", font=font(19, True), fill=LIME)
    wrapped(draw, "500 lecturas\n14 eventos\n10 ciclos de reúso\n2 zonas · 2 ESP32", (1400, 965, 300, 100), size=20, color=(240, 249, 245), spacing=7)
    path = DIAGRAMS / "Modelo_Datos_AgroVic_AquaCircular.png"
    image.save(path, quality=95)
    return path


def flow_diagram():
    image = Image.new("RGB", (1700, 900), PAPER)
    draw = ImageDraw.Draw(image)
    draw.text((75, 55), "Decisión de riego y reutilización", font=font(48, True), fill=INK)
    items = [
        ("1", "LECTURA", "Humedad, niveles,\ncaudal y calidad", FOREST_DARK),
        ("2", "VALIDACIÓN", "Rangos, vigencia,\ntoken e idempotencia", FOREST),
        ("3", "DECISIÓN", "¿Nivel y calidad de\nrecuperada son aptos?", (19, 124, 111)),
        ("4A", "REÚSO", "Prioriza estanque\nrecuperado", (31, 151, 132)),
        ("4B", "RESPALDO", "Usa fuente principal\ny genera alerta", (164, 105, 45)),
    ]
    coords = [(70, 260), (390, 260), (710, 260), (1100, 175), (1100, 500)]
    for (number, title, body, color), (x, y) in zip(items, coords):
        draw.rounded_rectangle((x, y, x + 260, y + 190), radius=22, fill=color)
        draw.text((x + 22, y + 18), number, font=font(18, True), fill=LIME)
        draw.text((x + 22, y + 55), title, font=font(20, True), fill=(255, 255, 255))
        wrapped(draw, body, (x + 22, y + 100, 215, 70), size=18, color=(230, 244, 238))
    arrow(draw, (330, 355), (390, 355))
    arrow(draw, (650, 355), (710, 355))
    draw.line([(970, 355), (1035, 355), (1035, 270), (1100, 270)], fill=AQUA, width=6)
    draw.polygon([(1100, 270), (1082, 260), (1082, 280)], fill=AQUA)
    draw.line([(970, 355), (1035, 355), (1035, 595), (1100, 595)], fill=(200, 132, 65), width=6)
    draw.polygon([(1100, 595), (1082, 585), (1082, 605)], fill=(200, 132, 65))
    draw.text((1000, 215), "SÍ", font=font(18, True), fill=FOREST)
    draw.text((990, 555), "NO", font=font(18, True), fill=(164, 105, 45))
    draw.rounded_rectangle((70, 720, 1360, 820), radius=18, fill=(226, 245, 238))
    draw.text((105, 750), "En cualquier rama: tiempo máximo 12 min, detención segura y auditoría.", font=font(24, True), fill=FOREST)
    path = DIAGRAMS / "Flujo_Decision_AgroVic_AquaCircular.png"
    image.save(path, quality=95)
    return path


def set_docx_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(8.5)


def docx_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_docx_cell_text(table.rows[0].cells[index], header, True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_docx_cell_text(cells[index], value)
    return table


def page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_docx(document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = DocxColor(15, 102, 92)
    header = section.header.paragraphs[0]
    header.text = "AGROVIC AQUACIRCULAR  |  INFORME FINAL"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = DocxColor(85, 112, 108)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Sebastián Tapia Véliz  ·  ")
    page_number(footer)


def add_picture_if(document, path, width=6.55):
    if Path(path).exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(width))


def generate_report(summary, diagrams):
    document = Document()
    configure_docx(document)
    document.core_properties.author = "Sebastián Tapia Véliz"
    document.core_properties.title = "Informe Final AgroVic AquaCircular"

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("\n\nIPLACEX\n").bold = True
    p.runs[-1].font.size = Pt(20)
    p.runs[-1].font.color.rgb = DocxColor(15, 102, 92)
    p.add_run("\nPROYECTO DE TÍTULO\n").bold = True
    p.runs[-1].font.size = Pt(14)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AgroVic AquaCircular")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = DocxColor(12, 55, 51)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Sistema web inteligente para el monitoreo, riego automatizado y reutilización controlada del agua mediante Arduino e IoT"
    ).font.size = Pt(15)
    document.add_paragraph("\n")
    cover = [
        "Estudiante: Sebastián Tapia Véliz",
        "Carrera: Técnico en Informática",
        "Asignatura: Proyecto de Título",
        "Docente: Aida Villamar Gallardo",
        "Caso de estudio: Agrícola Vicuña García Huidobro – AgroVic",
        "Fecha: septiembre de 2026",
    ]
    for line in cover:
        paragraph = document.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("\n")
    note = document.add_paragraph(
        "Proyecto académico. La solución se valida en un entorno reproducible con datos de prueba, firmware de referencia y simulador; no constituye una implementación oficial en terreno."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("Resumen ejecutivo", level=1)
    document.add_paragraph(
        "AgroVic AquaCircular responde a la falta de una visión consolidada del uso de agua, la pérdida de excedentes recuperables y el riesgo de reutilizar agua sin controles de nivel y calidad. La entrega reconstruye el prototipo como una aplicación funcional, reproducible y trazable. Integra una API Flask, MongoDB, dos perfiles, tres CRUD administrativos (usuarios, zonas y dispositivos), telemetría simulada de ESP32, reglas versionadas, alertas, control manual y reportes."
    )
    document.add_paragraph(
        f"La ejecución adjunta registró {summary['executed']} de 16 casos y {summary['passed']} aprobaciones, equivalentes a {summary['pass_rate']}%. La base contiene 500 lecturas y 14 ciclos cerrados: 197 L reutilizados de 294 L consumidos, lo que representa 67,01% de sustitución de agua principal en el conjunto de prueba. Estas cifras describen el dataset académico y no una medición productiva del predio."
    )
    document.add_paragraph(
        "El principal cambio frente a versiones anteriores es el cierre de trazabilidad: los casos CP01–CP16 ya no son marcadores pendientes. Cada evidencia EV se deriva de aserciones ejecutadas y conserva entrada, resultado observado y requisito relacionado. Además, se separa con claridad la validación de software de la futura validación física."
    )

    document.add_heading("Contenido", level=1)
    for item in [
        "1. Antecedentes y situación actual", "2. Objetivos y alcance", "3. Requisitos",
        "4. Metodología y correcciones", "5. Arquitectura", "6. Modelo de datos",
        "7. Implementación", "8. Seguridad", "9. Pruebas y evidencias",
        "10. Resultados y factibilidad", "11. Conclusiones", "Referencias", "Anexos",
    ]:
        document.add_paragraph(item)

    document.add_page_break()
    document.add_heading("1. Antecedentes y situación actual", level=1)
    document.add_paragraph(
        "El caso de estudio considera un piloto agrícola con dos zonas de cultivo. La situación inicial combina registros dispersos, decisiones de riego basadas en observación puntual, pérdida de lluvia o excedentes limpios y ausencia de una regla verificable para elegir entre agua recuperada y fuente principal. El efecto no es solo ambiental: también dificulta explicar consumos, detectar fallas y demostrar por qué se abrió una válvula."
    )
    document.add_paragraph(
        "Los actores son el Administrador, que configura usuarios, zonas, equipos y umbrales; el Operador, que supervisa y ejecuta acciones autorizadas; y el nodo ESP32, que reporta telemetría. La persona responsable del cultivo recibe información, pero la decisión automática siempre se restringe por calidad, nivel, vigencia y tiempo máximo."
    )
    document.add_heading("1.1 Problema central", level=2)
    document.add_paragraph(
        "No existe un mecanismo único y auditable que mida el balance hídrico, priorice el agua recuperada cuando sea apta y conduzca el riego a un estado seguro cuando una condición deja de cumplirse."
    )
    document.add_heading("1.2 Alternativa propuesta", level=2)
    document.add_paragraph(
        "Se propone una arquitectura de borde y servidor. El ESP32 conserva la detención local, mientras la aplicación valida lecturas, decide la fuente permitida, registra eventos y presenta información. Esta separación evita que una caída de la interfaz deje una bomba funcionando sin límite."
    )

    document.add_heading("2. Objetivos y alcance", level=1)
    document.add_paragraph(
        "El objetivo general es diseñar e implementar un sistema web inteligente para monitorear el riego y apoyar la reutilización controlada del agua mediante Arduino e IoT."
    )
    docx_table(document, ["Objetivo", "Criterio verificable", "Resultado"], OBJECTIVES)
    document.add_heading("2.1 Alcance incluido", level=2)
    document.add_paragraph(
        "Aplicación web responsive, perfiles Administrador y Operador, tres CRUD completos de administración, API de telemetría, dos zonas, reglas de riego, MongoDB exportable, firmware, simulador, pruebas, evidencias, documentación y paquete de entrega."
    )
    document.add_heading("2.2 Exclusiones y límites", level=2)
    document.add_paragraph(
        "No se incluyen instalación eléctrica o hidráulica definitiva, calibración certificada, conectividad rural contratada, integración con equipos productivos ni mediciones oficiales del predio. Estos elementos requieren trabajo en terreno, responsables de seguridad y validación agronómica."
    )

    document.add_page_break()
    document.add_heading("3. Requisitos del sistema", level=1)
    document.add_heading("3.1 Requisitos funcionales", level=2)
    docx_table(document, ["ID", "Descripción"], RF)
    document.add_heading("3.2 Requisitos no funcionales", level=2)
    docx_table(document, ["ID", "Descripción"], RNF)
    document.add_heading("3.3 Matriz de perfiles", level=2)
    docx_table(
        document,
        ["Acción", "Administrador", "Operador"],
        [
            ("Consultar panel, lecturas y balance", "Sí", "Sí"),
            ("Ejecutar control manual autorizado", "Sí", "Sí"),
            ("Confirmar alertas", "Sí", "Sí"),
            ("Cerrar alerta crítica", "Sí", "No"),
            ("Gestionar usuarios, zonas y dispositivos", "Sí", "No"),
            ("Versionar reglas y prioridad de fuente", "Sí", "No"),
            ("Consultar auditoría y exportar", "Sí", "Exportar"),
        ],
    )

    document.add_heading("4. Metodología y correcciones incorporadas", level=1)
    document.add_paragraph(
        "La reconstrucción se organizó en requisitos, implementación, datos, pruebas y documentación. Se utilizó una base determinista para que cada ejecución pueda compararse. El pipeline instala dependencias, inicia MongoDB, ejecuta pytest, vuelve a sembrar la demostración, captura la interfaz y construye documentos y ZIP."
    )
    docx_table(
        document,
        ["Hallazgo previo", "Corrección aplicada", "Evidencia"],
        [
            ("Casos CP01–CP16 sin ejecución", "Pruebas automatizadas independientes", "test_results.json y EV01–EV16"),
            ("Enlaces y entregables pendientes", "Repositorio público y paquete versionado", "README, workflow y Release"),
            ("Resultados declarativos", "Métricas calculadas desde MongoDB", "CP15 y balance CSV"),
            ("Alcance de hardware ambiguo", "Separación entre simulador y terreno", "Firmware, simulador y límite documentado"),
            ("Perfiles sin demostración", "Autorización en servidor y vistas diferenciadas", "CP03, CP13 y capturas"),
            ("Modelo sin exportables", "Extended JSON por colección", "database/exports"),
        ],
    )

    document.add_page_break()
    document.add_heading("5. Arquitectura de solución", level=1)
    document.add_paragraph(
        "La capa de borde reúne sensores y actuadores. La API recibe JSON autenticado, valida rangos y evita duplicados. MongoDB conserva series, maestros y eventos. La interfaz consulta el mismo dominio de reglas, evitando diferencias entre lo mostrado y lo ejecutado."
    )
    add_picture_if(document, diagrams["architecture"])
    document.add_heading("5.1 Flujo de decisión", level=2)
    document.add_paragraph(
        "Una lectura válida se compara con la regla activa. El agua recuperada solo se autoriza cuando la lectura de calidad es válida, la turbidez y conductividad están bajo máximos y el nivel supera el mínimo. En caso contrario se selecciona la fuente principal y se crea una alerta. Un watchdog detiene eventos que exceden doce minutos y marca desconexiones después de diez minutos sin telemetría."
    )
    add_picture_if(document, diagrams["flow"])
    document.add_heading("5.2 Factibilidad técnica", level=2)
    document.add_paragraph(
        "La pila utiliza componentes ampliamente disponibles: Python, Flask, MongoDB, Docker y ESP32. El diseño monolítico modular es adecuado para el piloto y reduce operación. La escalabilidad se obtiene agregando zonas y dispositivos sin cambiar la estructura base."
    )

    document.add_page_break()
    document.add_heading("6. Modelo de datos MongoDB", level=1)
    document.add_paragraph(
        "Se eligió MongoDB por la naturaleza variable de la telemetría y la necesidad de conservar documentos con metadatos de zona y dispositivo. lecturas_sensor se crea como colección time-series con timestamp y meta. Las restantes colecciones utilizan validación JSON Schema e índices."
    )
    add_picture_if(document, diagrams["data"])
    docx_table(document, ["Colección", "Responsabilidad"], COLLECTIONS)
    document.add_heading("6.1 Exportación e importación", level=2)
    document.add_paragraph(
        "El script de siembra exporta un arreglo Extended JSON por colección. La entrega incluye instrucciones de mongoimport. No se almacenan datos personales reales: correos, contraseñas y token pertenecen exclusivamente a la demostración generada."
    )

    document.add_heading("7. Implementación", level=1)
    document.add_heading("7.1 Aplicación web y formularios", level=2)
    document.add_paragraph(
        "El panel presenta sustitución, consumo, captura, ciclos, humedad, niveles, calidad, caudal, fuente autorizada, alertas y eventos. La administración incorpora tres CRUD verificables: usuarios, zonas y dispositivos; además se gestionan reglas de riego y acciones operativas. Cada CRUD permite crear, consultar, modificar y realizar una baja lógica con validaciones en el servidor. Las operaciones conservan la trazabilidad histórica."
    )
    for screenshot in [
        EVIDENCE / "UI01_panel_administrador.png",
        EVIDENCE / "UI02_formularios_administracion.png",
        EVIDENCE / "UI05_panel_operador_movil.png",
    ]:
        add_picture_if(document, screenshot)
    document.add_heading("7.1.1 Gestión completa de los tres CRUD", level=2)
    document.add_paragraph(
        "Los módulos Usuarios y Operadores, Zonas de Cultivo y Dispositivos ESP32 permiten crear, "
        "consultar, modificar y eliminar mediante desactivación lógica. Esta última operación conserva "
        "los documentos y su vínculo con auditoría, mediciones y eventos. La interfaz impide que un "
        "administrador elimine su propia cuenta y rechaza la desactivación de una zona con riego activo. "
        "Los dispositivos asociados a una zona desactivada también quedan inactivos."
    )
    document.add_paragraph(
        "La evidencia automática ejecuta altas, cambios y bajas reales en la base de demostración "
        "de GitHub Actions y verifica que el estado persiste. No se utilizan datos de producción; "
        "el token inicial del ESP32 no se almacena en claro ni aparece en capturas."
    )
    for screenshot in [
        EVIDENCE / "UI06_crud_usuarios_crear_leer.png",
        EVIDENCE / "UI08_crud_usuarios_desactivado.png",
        EVIDENCE / "UI09_crud_zonas_crear_leer.png",
        EVIDENCE / "UI11_crud_zonas_desactivada.png",
        EVIDENCE / "UI12_crud_dispositivos_crear_leer.png",
        EVIDENCE / "UI14_crud_dispositivos_desactivado.png",
    ]:
        add_picture_if(document, screenshot, 6.1)
    document.add_heading("7.2 API", level=2)
    docx_table(document, ["Método", "Ruta", "Propósito"], API)
    document.add_heading("7.3 Dataset de demostración", level=2)
    document.add_paragraph(
        "Las 500 lecturas se distribuyen en dos zonas con intervalos de cinco minutos. Los 14 eventos cerrados suman 294 L; 10 eventos utilizan 197 L recuperados y cuatro utilizan 97 L principales. La captura acumulada del conjunto es 240 L."
    )

    document.add_heading("8. Seguridad y continuidad", level=1)
    document.add_paragraph(
        "Las contraseñas se guardan con bcrypt. La sesión se protege con secreto generado o suministrado por entorno. Los formularios usan CSRF; las rutas aplican permisos en servidor. El token del dispositivo nunca se guarda en claro en MongoDB. La recuperación usa hash, expiración y un solo uso. Se agregan cabeceras CSP, nosniff, frame deny y política de permisos."
    )
    document.add_paragraph(
        "HTTPS se delega a un proxy inverso en producción. Esta entrega local no finge un certificado público. Para continuidad, el firmware detiene la bomba ante pérdida de conectividad y el servidor dispone de watchdog para tiempo máximo y desconexión."
    )

    document.add_page_break()
    document.add_heading("9. Pruebas y evidencias", level=1)
    document.add_paragraph(
        f"Resultado de la ejecución: {summary['passed']}/{summary['total']} casos aprobados ({summary['pass_rate']}%). Criterio mínimo: 15/16 y casos críticos aprobados. Las evidencias que siguen se generaron desde los resultados observados."
    )
    rows = [
        (item["case"], item["title"], item["status"], item["observed"], ", ".join(item["requirements"]))
        for item in summary["results"]
    ]
    docx_table(document, ["Caso", "Propósito", "Estado", "Resultado observado", "Trazabilidad"], rows)
    document.add_heading("9.1 Evidencias EV01–EV16", level=2)
    for item in summary["results"]:
        document.add_heading(f"{item['evidence']} · {item['case']} · {item['title']}", level=3)
        document.add_paragraph(f"Entrada o acción: {item['inputs']}")
        document.add_paragraph(f"Resultado: {item['observed']}")
        evidence_path = EVIDENCE / f"{item['evidence']}_{item['case']}.png"
        add_picture_if(document, evidence_path, 6.25)

    document.add_heading("10. Resultados, factibilidad y riesgos", level=1)
    document.add_paragraph(
        "El resultado principal es una cadena demostrable desde la medición hasta el reporte. El indicador de sustitución de 67,01% verifica el cálculo y la utilidad del tablero, pero no debe extrapolarse a una temporada real. Para hacerlo se requiere calibración, muestreo representativo y comparación con registros de terreno."
    )
    docx_table(
        document,
        ["Dimensión", "Evaluación", "Condición"],
        [
            ("Técnica", "Factible para piloto", "Calibrar sensores y adaptar actuadores"),
            ("Operacional", "Flujo simple por roles", "Capacitar y definir responsables"),
            ("Económica", "Componentes de bajo costo relativo", "Cotizar instalación y mantenimiento"),
            ("Legal/seguridad", "Datos de prueba sin información personal", "Revisar normativa de agua y seguridad eléctrica"),
            ("Ambiental", "Permite medir sustitución", "Validar calidad según cultivo y uso"),
        ],
    )
    document.add_heading("10.1 Riesgos y mitigaciones", level=2)
    docx_table(
        document,
        ["Riesgo", "Impacto", "Mitigación"],
        [
            ("Sensor descalibrado", "Decisión incorrecta", "Calibración, rangos y comparación manual"),
            ("Pérdida de red", "Telemetría ausente", "Detención local y alerta por watchdog"),
            ("Agua fuera de calidad", "Daño al cultivo o equipo", "Bloqueo de reúso y fuente principal"),
            ("Acceso indebido", "Cambio de reglas o accionamiento", "bcrypt, roles, CSRF, HTTPS y auditoría"),
            ("Sobreestimación de ahorro", "Conclusión no válida", "Separar dataset académico y medición productiva"),
        ],
    )

    document.add_heading("11. Conclusiones", level=1)
    document.add_paragraph(
        "La entrega demuestra que el problema puede abordarse con un sistema pequeño pero completo: mide, valida, decide, registra y explica. La prioridad de agua recuperada no es incondicional; depende de calidad y nivel. Esta regla, junto con el watchdog y los perfiles, convierte una interfaz de monitoreo en un prototipo operacionalmente responsable."
    )
    document.add_paragraph(
        "Los objetivos de software, datos, interfaz y pruebas se cumplen en el entorno reproducible. El objetivo de integración se cumple a nivel de protocolo, firmware y simulación. Su cierre productivo exige hardware conectado, calibración e instalación en terreno, por lo que se mantiene como paso futuro y no como evidencia fabricada."
    )
    document.add_paragraph(
        "La siguiente etapa recomendada es un piloto físico acotado con una sola válvula, registro paralelo manual, criterios agronómicos aprobados y revisión semanal de falsos positivos. Solo después corresponde ampliar zonas o automatizar decisiones sin supervisión."
    )

    document.add_heading("Referencias", level=1)
    for reference in [
        "Arduino. (s. f.). ArduinoJson documentation. https://arduinojson.org/",
        "Espressif Systems. (2026). ESP32 series datasheet. https://www.espressif.com/",
        "MongoDB, Inc. (2026). MongoDB manual. https://www.mongodb.com/docs/manual/",
        "OWASP Foundation. (2021). OWASP Top 10. https://owasp.org/www-project-top-ten/",
        "Pallets Projects. (2026). Flask documentation. https://flask.palletsprojects.com/",
        "Python Software Foundation. (2026). Python 3 documentation. https://docs.python.org/3/",
        "W3C. (2023). Web Content Accessibility Guidelines 2.2. https://www.w3.org/TR/WCAG22/",
    ]:
        document.add_paragraph(reference)

    document.add_heading("Anexos", level=1)
    document.add_paragraph("A. Repositorio público: https://github.com/seba32111-svg/agrovic-aquacircular")
    document.add_paragraph("B. Archivos JSON exportables: database/exports/")
    document.add_paragraph("C. Resultados de prueba: artifacts/test_results.json")
    document.add_paragraph("D. Evidencias visuales: artifacts/evidence/")
    document.add_paragraph("E. Firmware y simulador: hardware/ y simulator/")

    output = DIST / "Informe_Final_AgroVic_AquaCircular_Sebastian_Tapia.docx"
    document.save(output)
    return output


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=28, leading=33, textColor=colors.HexColor("#0f665c"), alignment=TA_CENTER, spaceAfter=18))
    styles.add(ParagraphStyle(name="H1Aqua", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.HexColor("#0f665c"), spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle(name="H2Aqua", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, textColor=colors.HexColor("#123d35"), spaceBefore=9, spaceAfter=6))
    styles.add(ParagraphStyle(name="BodyAqua", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13.5, textColor=colors.HexColor("#10272b"), spaceAfter=7))
    styles.add(ParagraphStyle(name="SmallAqua", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=10))
    return styles


def pdf_table(rows, widths=None, repeat=1):
    data = [[Paragraph(str(cell), getSampleStyleSheet()["BodyText"]) for cell in row] for row in rows]
    table = LongTable(data, colWidths=widths, repeatRows=repeat, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f665c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.2),
        ("LEADING", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#cdded8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f7f3")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def generate_report_pdf(summary, diagrams):
    output = DIST / "Informe_Final_AgroVic_AquaCircular_Sebastian_Tapia.pdf"
    styles = pdf_styles()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#b8d4cb"))
        canvas.line(2 * cm, 1.35 * cm, 19 * cm, 1.35 * cm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#526f69"))
        canvas.drawString(2 * cm, .9 * cm, "AgroVic AquaCircular · Sebastián Tapia Véliz")
        canvas.drawRightString(19 * cm, .9 * cm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=1.7*cm, leftMargin=1.7*cm, topMargin=1.7*cm, bottomMargin=1.7*cm)
    story = [
        Spacer(1, 2.2*cm),
        Paragraph("IPLACEX · PROYECTO DE TÍTULO", styles["H2Aqua"]),
        Spacer(1, .8*cm),
        Paragraph("AgroVic AquaCircular", styles["CoverTitle"]),
        Paragraph("Sistema web inteligente para monitoreo, riego automatizado y reutilización controlada del agua mediante Arduino e IoT", ParagraphStyle("Sub", parent=styles["BodyAqua"], fontSize=14, leading=19, alignment=TA_CENTER)),
        Spacer(1, 2*cm),
        Paragraph("Sebastián Tapia Véliz · Técnico en Informática", ParagraphStyle("Center", parent=styles["BodyAqua"], alignment=TA_CENTER, fontSize=11)),
        Paragraph("Docente: Aida Villamar Gallardo · septiembre de 2026", ParagraphStyle("Center2", parent=styles["BodyAqua"], alignment=TA_CENTER)),
        Spacer(1, 2*cm),
        Paragraph("Proyecto académico validado con datos de prueba, firmware de referencia y simulador reproducible. No constituye una implementación oficial en terreno.", ParagraphStyle("Note", parent=styles["BodyAqua"], alignment=TA_CENTER, textColor=colors.HexColor("#526f69"))),
        PageBreak(),
        Paragraph("Resumen ejecutivo", styles["H1Aqua"]),
        Paragraph(f"La entrega integra aplicación funcional, MongoDB, perfiles Administrador y Operador, formularios, API de ESP32, alertas, reportes y auditoría. La ejecución obtuvo {summary['passed']}/{summary['total']} casos aprobados ({summary['pass_rate']}%).", styles["BodyAqua"]),
        Paragraph("El conjunto demostrativo contiene 500 lecturas y 14 ciclos. Se reutilizan 197 L de 294 L, equivalentes a 67,01% de sustitución. La cifra valida el cálculo del sistema y no se extrapola a producción.", styles["BodyAqua"]),
        Paragraph("1. Problema y solución", styles["H1Aqua"]),
        Paragraph("La situación inicial carece de una vista consolidada del consumo, pierde excedentes recuperables y no documenta con qué criterio se autoriza el reúso. La solución separa seguridad local, API, datos y experiencia web.", styles["BodyAqua"]),
        PdfImage(str(diagrams["architecture"]), width=17*cm, height=9.9*cm),
        Paragraph("2. Objetivos", styles["H1Aqua"]),
        pdf_table([["ID", "Criterio", "Resultado"]] + [list(row) for row in OBJECTIVES], [1.2*cm, 10.5*cm, 5.3*cm]),
        PageBreak(),
        Paragraph("3. Requisitos funcionales", styles["H1Aqua"]),
        pdf_table([["ID", "Descripción"]] + [list(row) for row in RF], [1.5*cm, 15.5*cm]),
        Paragraph("4. Requisitos no funcionales", styles["H1Aqua"]),
        pdf_table([["ID", "Descripción"]] + [list(row) for row in RNF], [1.5*cm, 15.5*cm]),
        PageBreak(),
        Paragraph("5. Arquitectura y seguridad", styles["H1Aqua"]),
        Paragraph("El servidor valida rangos, vigencia e idempotencia. La selección de fuente exige calidad y nivel aptos. bcrypt, permisos por rol, CSRF, hash de token, recuperación temporal, cabeceras defensivas y auditoría forman la línea base.", styles["BodyAqua"]),
        PdfImage(str(diagrams["flow"]), width=17*cm, height=9*cm),
        Paragraph("6. Modelo MongoDB", styles["H1Aqua"]),
        PdfImage(str(diagrams["data"]), width=17*cm, height=11.1*cm),
        PageBreak(),
        Paragraph("7. Implementación", styles["H1Aqua"]),
        Paragraph("La administración dispone de tres CRUD completos (usuarios, zonas y dispositivos), con bajas lógicas y control por roles. Además, el panel gestiona reglas, balance, sensores, fuentes, alertas e historial.", styles["BodyAqua"]),
    ]
    for screenshot in [EVIDENCE / "UI01_panel_administrador.png", EVIDENCE / "UI02_formularios_administracion.png", EVIDENCE / "UI05_panel_operador_movil.png"]:
        if screenshot.exists():
            with Image.open(screenshot) as img:
                ratio = img.height / img.width
            width = 16.5*cm
            height = min(20*cm, width * ratio)
            story += [PdfImage(str(screenshot), width=width, height=height), Spacer(1, .4*cm)]
    story += [
        PageBreak(),
        Paragraph("7.1 Evidencias de mantenimiento CRUD", styles["H1Aqua"]),
        Paragraph("Las operaciones de alta, consulta, edición y desactivación se verifican en MongoDB de prueba. Los tres módulos conservan historial y restringen cambios al administrador.", styles["BodyAqua"]),
        Paragraph("8. Pruebas", styles["H1Aqua"]),
        Paragraph(f"Resultado: {summary['passed']} aprobados de {summary['total']} ({summary['pass_rate']}%). Cada EV se construye desde el resultado real de pytest.", styles["BodyAqua"]),
        pdf_table(
            [["Caso", "Propósito", "Estado", "Resultado"]] +
            [[item["case"], item["title"], item["status"], item["observed"]] for item in summary["results"]],
            [1.3*cm, 4.2*cm, 2.1*cm, 9.4*cm],
        ),
        PageBreak(),
        Paragraph("9. Conclusiones", styles["H1Aqua"]),
        Paragraph("El prototipo completa la cadena medir-validar-decidir-registrar-explicar. La reutilización se autoriza solo bajo condiciones aptas y cualquier fallo relevante conduce a una alternativa segura.", styles["BodyAqua"]),
        Paragraph("La integración física debe validarse en terreno con calibración, seguridad eléctrica e hidráulica y criterio agronómico. La presente evidencia cubre software, base de datos, protocolo, firmware de referencia y simulación.", styles["BodyAqua"]),
        Paragraph("Referencias", styles["H1Aqua"]),
        Paragraph("MongoDB Manual (2026); Flask Documentation (2026); OWASP Top 10 (2021); Espressif ESP32 Datasheet (2026); WCAG 2.2 (2023). Enlaces completos disponibles en docs/FUENTES.md.", styles["BodyAqua"]),
    ]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output


def add_text(slide, x, y, w, h, text, size=22, color=(255,255,255), bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.font.name = "Aptos"
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def rect(slide, x, y, w, h, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, PInches(x), PInches(y), PInches(w), PInches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(*(line or fill))
    return shape


def slide_base(prs, number, title, kicker="AGROVIC AQUACIRCULAR", dark=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = FOREST_DARK if dark else PAPER
    rect(slide, 0, 0, 13.333, 7.5, background, False)
    rect(slide, 0, 0, .12, 7.5, AQUA, False)
    title_color = (255,255,255) if dark else INK
    muted = (151,198,188) if dark else MUTED
    add_text(slide, .62, .34, 6.8, .3, kicker, 9, muted, True)
    add_text(slide, .62, .76, 11.5, .72, title, 28, title_color, True)
    add_text(slide, 10.85, .33, 1.75, .3, "IPLACEX", 11, LIME if dark else FOREST, True, PP_ALIGN.RIGHT)
    add_text(slide, 12.15, 7.08, .55, .2, f"{number:02d}", 9, muted, True, PP_ALIGN.RIGHT)
    return slide


def bullet_block(slide, x, y, w, lines, dark=True, size=18):
    color = (235,246,241) if dark else INK
    for index, line in enumerate(lines):
        rect(slide, x, y + index*.72 + .16, .12, .12, AQUA, True)
        add_text(slide, x+.28, y + index*.72, w-.28, .55, line, size, color)


def picture(slide, path, x, y, w, h=None):
    if Path(path).exists():
        if h:
            slide.shapes.add_picture(str(path), PInches(x), PInches(y), PInches(w), PInches(h))
        else:
            slide.shapes.add_picture(str(path), PInches(x), PInches(y), width=PInches(w))


def metric_card(slide, x, y, w, label, value, note="", featured=False):
    fill = (18,126,113) if featured else (244,248,245)
    color = (255,255,255) if featured else INK
    shape = rect(slide, x, y, w, 1.38, fill, True)
    add_text(slide, x+.2, y+.17, w-.4, .22, label.upper(), 8, LIME if featured else MUTED, True)
    add_text(slide, x+.2, y+.52, w-.4, .46, value, 25, color, True)
    if note:
        add_text(slide, x+.2, y+1.04, w-.4, .2, note, 8, (198,235,225) if featured else MUTED)


def generate_presentation(summary, diagrams):
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    prs.core_properties.author = "Sebastián Tapia Véliz"
    prs.core_properties.title = "AgroVic AquaCircular — Presentación final"

    slide = slide_base(prs, 1, "", "IPLACEX · PROYECTO DE TÍTULO")
    add_text(slide, .75, 1.35, 7.2, 1.2, "AgroVic\nAquaCircular", 45, (255,255,255), True)
    add_text(slide, .78, 3.85, 6.6, 1.0, "Monitoreo, riego y reutilización controlada del agua mediante ESP32 e IoT", 20, (197,230,221))
    rect(slide, 9.1, 1.35, 2.7, 3.8, (15,91,80), True)
    add_text(slide, 9.7, 2.02, 1.5, .5, "H₂O", 34, LIME, True, PP_ALIGN.CENTER)
    add_text(slide, 9.52, 3.0, 1.9, 1.2, "MEDIR\nDECIDIR\nREUTILIZAR", 15, (226,245,238), True, PP_ALIGN.CENTER)
    add_text(slide, .78, 6.45, 6.5, .35, "Sebastián Tapia Véliz · Técnico en Informática · 2026", 10, (151,198,188))

    slide = slide_base(prs, 2, "Identificación del proyecto", dark=False)
    add_text(slide, .65, 1.7, 6.0, 1.4, "Un sistema pequeño,\nuna decisión hídrica trazable.", 30, FOREST_DARK, True)
    bullet_block(slide, .7, 3.55, 5.7, [
        "Estudiante: Sebastián Tapia Véliz",
        "Carrera: Técnico en Informática",
        "Docente: Aida Villamar Gallardo",
        "Caso académico: AgroVic, Vicuña",
    ], dark=False, size=15)
    metric_card(slide, 8.0, 1.85, 2.0, "Zonas", "2", "piloto")
    metric_card(slide, 10.2, 1.85, 2.0, "Perfiles", "2", "Admin / Operador", True)
    metric_card(slide, 8.0, 3.55, 2.0, "Lecturas", "500", "cada 5 min")
    metric_card(slide, 10.2, 3.55, 2.0, "Pruebas", "16", "CP01–CP16")

    slide = slide_base(prs, 3, "Introducción")
    add_text(slide, .68, 1.7, 5.3, 1.6, "El agua recuperada solo aporta valor cuando se puede medir, validar y explicar.", 27, (255,255,255), True)
    bullet_block(slide, .75, 3.75, 5.4, ["Telemetría por zona", "Reglas de nivel y calidad", "Decisiones con trazabilidad"], True, 17)
    rect(slide, 7.2, 1.7, 4.9, 4.7, (18,80,71), True)
    add_text(slide, 7.65, 2.15, 3.9, .3, "PREGUNTA RECTORA", 10, LIME, True)
    add_text(slide, 7.65, 2.75, 3.8, 2.4, "¿Cómo priorizar el reúso sin comprometer la seguridad del riego?", 28, (239,249,245), True)
    add_text(slide, 7.65, 5.55, 3.8, .45, "Respuesta: condición + respaldo + auditoría", 12, (158,207,197))

    slide = slide_base(prs, 4, "Situación actual", dark=False)
    cards = [
        ("01", "Información dispersa", "No existe un balance único por zona y fuente."),
        ("02", "Excedente perdido", "Lluvia o agua limpia recuperable no se cuantifica."),
        ("03", "Reúso riesgoso", "Sin nivel y calidad, la prioridad puede ser incorrecta."),
        ("04", "Baja trazabilidad", "No queda claro quién actuó y por qué."),
    ]
    for index, (n, title, body) in enumerate(cards):
        x = .7 + (index%2)*6.05
        y = 1.65 + (index//2)*2.3
        rect(slide, x, y, 5.55, 1.9, (255,255,255), True, (210,225,219))
        add_text(slide, x+.24, y+.22, .5, .25, n, 10, AQUA, True)
        add_text(slide, x+.85, y+.2, 4.2, .35, title, 18, FOREST_DARK, True)
        add_text(slide, x+.85, y+.72, 4.2, .75, body, 14, MUTED)

    slide = slide_base(prs, 5, "Alternativa de solución")
    picture(slide, diagrams["flow"], .55, 1.55, 8.0, 4.25)
    rect(slide, 9.0, 1.62, 3.45, 4.85, (18,80,71), True)
    add_text(slide, 9.35, 1.98, 2.7, .28, "PRINCIPIO", 9, LIME, True)
    add_text(slide, 9.35, 2.45, 2.7, 1.4, "Recuperada\nsi es apta.", 27, (255,255,255), True)
    add_text(slide, 9.35, 4.02, 2.7, 1.3, "Principal si no.\nSiempre registrar.", 18, (193,228,219), True)
    add_text(slide, 9.35, 5.75, 2.7, .35, "Tiempo máximo: 12 min", 11, LIME, True)

    slide = slide_base(prs, 6, "Factibilidad de la alternativa", dark=False)
    feasibility = [
        ("TÉCNICA", "Flask · MongoDB · ESP32", "Componentes disponibles y arquitectura reproducible."),
        ("OPERACIONAL", "Dos perfiles claros", "Acciones críticas restringidas y auditadas."),
        ("ECONÓMICA", "Piloto incremental", "Escala desde dos zonas sin rediseñar el dominio."),
        ("AMBIENTAL", "Balance medible", "Distingue agua capturada, reutilizada y principal."),
    ]
    for index, (label, value, detail) in enumerate(feasibility):
        x = .7 + index*3.08
        rect(slide, x, 1.75, 2.75, 4.65, (255,255,255), True, (205,222,216))
        add_text(slide, x+.25, 2.08, 2.2, .25, label, 9, AQUA, True)
        add_text(slide, x+.25, 2.65, 2.2, 1.0, value, 20, FOREST_DARK, True)
        add_text(slide, x+.25, 4.18, 2.2, 1.2, detail, 13, MUTED)

    slide = slide_base(prs, 7, "Descripción del sistema")
    picture(slide, diagrams["architecture"], .55, 1.48, 8.7, 5.08)
    bullet_block(slide, 9.65, 1.82, 2.8, [
        "Borde seguro",
        "API validada",
        "MongoDB",
        "Web responsive",
        "Exportación",
    ], True, 16)

    slide = slide_base(prs, 8, "Modelo y relaciones", dark=False)
    picture(slide, diagrams["data"], .55, 1.45, 9.6, 5.45)
    rect(slide, 10.45, 1.7, 2.25, 4.9, FOREST, True)
    add_text(slide, 10.78, 2.1, 1.6, .4, "10", 35, LIME, True)
    add_text(slide, 10.78, 2.62, 1.6, .5, "colecciones", 13, (230,246,240), True)
    add_text(slide, 10.78, 3.55, 1.55, 1.6, "Serie temporal\n+ maestros\n+ eventos\n+ auditoría", 14, (195,228,219))
    add_text(slide, 10.78, 5.75, 1.55, .5, "Extended JSON", 10, LIME, True)

    slide = slide_base(prs, 9, "Requisitos funcionales")
    groups = [
        ("ACCESO", "RF01–RF03", "Login, recuperación, roles y usuarios"),
        ("OPERACIÓN", "RF04–RF11", "Zonas, sensores, fuentes, riego y seguridad"),
        ("INFORMACIÓN", "RF12–RF15", "Panel, alertas, reportes y auditoría"),
    ]
    for index, (label, ids, detail) in enumerate(groups):
        x = .75 + index*4.12
        rect(slide, x, 1.8, 3.65, 4.55, (17,78,69), True)
        add_text(slide, x+.3, 2.12, 2.9, .25, label, 9, LIME, True)
        add_text(slide, x+.3, 2.7, 2.9, .55, ids, 27, (255,255,255), True)
        add_text(slide, x+.3, 3.7, 2.9, 1.4, detail, 17, (194,228,219))
        add_text(slide, x+.3, 5.7, 2.9, .3, f"{3 if index==0 else 8 if index==1 else 4} requisitos", 10, (145,197,186), True)

    slide = slide_base(prs, 10, "Requisitos no funcionales", dark=False)
    add_text(slide, .7, 1.62, 5.2, .7, "Seguridad y continuidad como comportamiento, no como anexo.", 24, FOREST_DARK, True)
    bullet_block(slide, .75, 2.75, 5.3, ["bcrypt y recuperación de un uso", "Roles, CSRF y validación", "Token de equipo almacenado como hash", "Watchdog y detención local"], False, 15)
    rect(slide, 7.0, 1.65, 5.25, 4.9, FOREST_DARK, True)
    metric_card(slide, 7.42, 2.1, 1.9, "Panel p95", "< 3 s", "criterio")
    metric_card(slide, 9.55, 2.1, 1.9, "Vista", "360 px", "mínimo", True)
    metric_card(slide, 7.42, 3.78, 1.9, "Meta", "99%", "disponibilidad")
    metric_card(slide, 9.55, 3.78, 1.9, "Envío", "≤ 5 min", "telemetría", True)
    add_text(slide, 7.45, 5.75, 3.9, .35, "HTTPS exigido para producción", 11, (173,215,205), True)

    slide = slide_base(prs, 11, "Implementación de datos")
    metric_card(slide, .72, 1.75, 2.65, "Lecturas", "500", "250 por zona", True)
    metric_card(slide, 3.62, 1.75, 2.65, "Ciclos", "14", "cerrados")
    metric_card(slide, 6.52, 1.75, 2.65, "Reúso", "10", "ciclos")
    metric_card(slide, 9.42, 1.75, 2.65, "Captura", "240 L", "dataset")
    add_text(slide, .75, 3.65, 5.1, .5, "Colección time-series", 21, (255,255,255), True)
    bullet_block(slide, .8, 4.25, 5.1, ["timestamp + meta", "índices por zona y equipo", "exportación por colección"], True, 15)
    add_text(slide, 7.0, 3.6, 4.8, .5, "Balance de prueba", 21, (255,255,255), True)
    add_text(slide, 7.0, 4.42, 4.8, 1.25, "197 L reutilizados\n÷ 294 L consumidos", 25, LIME, True)
    add_text(slide, 7.0, 5.92, 4.8, .4, "= 67,01% sustitución", 16, (191,229,218), True)

    slide = slide_base(prs, 12, "Administración: tres CRUD completos", dark=False)
    add_text(slide, .75, 1.52, 11.7, .45, "Crear · Consultar · Modificar · Desactivar / Reactivar", 20, FOREST_DARK, True)
    evidence_cases = [
        ("CRUD 01 · USUARIOS", "UI06_crud_usuarios_crear_leer.png"),
        ("CRUD 02 · ZONAS", "UI09_crud_zonas_crear_leer.png"),
        ("CRUD 03 · DISPOSITIVOS", "UI12_crud_dispositivos_crear_leer.png"),
    ]
    for index, (label, filename) in enumerate(evidence_cases):
        x = .55 + index * 4.24
        rect(slide, x, 2.25, 3.95, 3.57, (255,255,255), True, (210,225,219))
        add_text(slide, x+.24, 2.5, 3.45, .35, label, 12, FOREST, True)
        picture(slide, EVIDENCE / filename, x+.2, 3.1, 3.55, 1.78)
        add_text(slide, x+.24, 5.3, 3.43, .33, "Prueba real en CI · MongoDB aislado", 9, MUTED)
    add_text(slide, .8, 6.32, 11.75, .36, "Baja lógica, protección por roles, auditoría y preservación de datos históricos.", 13, FOREST_DARK, True)

    slide = slide_base(prs, 13, "Pruebas y evidencias")
    metric_card(slide, .72, 1.72, 2.55, "Aprobados", f"{summary['passed']}/{summary['total']}", "CP01–CP16", True)
    metric_card(slide, 3.52, 1.72, 2.55, "Cumplimiento", f"{summary['pass_rate']}%", "mínimo 93,75%")
    metric_card(slide, 6.32, 1.72, 2.55, "Evidencias", str(summary["executed"]), "EV generadas")
    metric_card(slide, 9.12, 1.72, 2.55, "Exportación", "CSV + JSON", "reproducible")
    picture(slide, EVIDENCE / "UI01_panel_administrador.png", .72, 3.48, 5.5, 3.1)
    picture(slide, EVIDENCE / "EV15_CP15.png", 6.55, 3.48, 5.9, 2.85)

    slide = slide_base(prs, 14, "Conclusiones", dark=False)
    add_text(slide, .72, 1.62, 6.4, 1.3, "Medir no basta.\nLa solución debe decidir y explicar.", 30, FOREST_DARK, True)
    conclusions = [
        ("01", "Reúso condicionado", "Nivel y calidad antes de priorizar."),
        ("02", "Seguridad en capas", "Servidor y watchdog local."),
        ("03", "Resultado verificable", "Datos, pruebas y evidencias reproducibles."),
    ]
    for index, (number, title, detail) in enumerate(conclusions):
        y = 3.38 + index*.95
        add_text(slide, .82, y, .4, .3, number, 10, AQUA, True)
        add_text(slide, 1.42, y-.03, 2.5, .35, title, 15, FOREST, True)
        add_text(slide, 4.2, y-.02, 3.1, .4, detail, 13, MUTED)
    rect(slide, 8.25, 1.65, 4.0, 4.95, FOREST_DARK, True)
    add_text(slide, 8.65, 2.05, 3.2, .25, "SIGUIENTE ETAPA", 9, LIME, True)
    add_text(slide, 8.65, 2.75, 3.05, 1.3, "Piloto físico\nacotado y calibrado", 25, (255,255,255), True)
    add_text(slide, 8.65, 4.45, 3.0, 1.15, "Una válvula · registro paralelo · validación agronómica", 15, (181,220,210))

    slide = slide_base(prs, 15, "", "IPLACEX · CIERRE")
    add_text(slide, .75, 1.55, 8.4, 1.2, "Agua circular.\nDecisiones claras.", 42, (255,255,255), True)
    add_text(slide, .78, 4.5, 6.8, .55, "Demostración funcional y preguntas", 20, LIME, True)
    add_text(slide, .78, 6.15, 6.8, .3, "Sebastián Tapia Véliz · AgroVic AquaCircular", 11, (151,198,188))
    rect(slide, 9.65, 1.6, 2.2, 3.8, (15,91,80), True)
    add_text(slide, 10.05, 2.35, 1.4, .5, "67,01%", 29, LIME, True, PP_ALIGN.CENTER)
    add_text(slide, 10.05, 3.08, 1.4, 1.0, "sustitución\nen datos de prueba", 13, (222,242,236), True, PP_ALIGN.CENTER)

    assert len(prs.slides) == 15
    output = DIST / "Presentacion_AgroVic_AquaCircular_Sebastian_Tapia.pptx"
    prs.save(output)
    return output


SCRIPT = [
    ("1. Portada", "0:45", "Buenos días. Soy Sebastián Tapia Véliz y presentaré AgroVic AquaCircular, un sistema académico para medir el riego, priorizar agua recuperada cuando sea apta y mantener trazabilidad. Primero explicaré el problema y el diseño; después mostraré la aplicación funcionando."),
    ("2. Identificación", "0:50", "El piloto considera dos zonas, dos perfiles y un nodo ESP32 por zona. La propuesta se desarrolló para el caso AgroVic en Vicuña, pero esta entrega no afirma una instalación oficial. Lo demostrable hoy es la solución de software, los datos, el firmware de referencia y la simulación."),
    ("3. Introducción", "1:00", "La pregunta central no es solo cuánto regar. Es cómo reutilizar sin convertir el ahorro en un riesgo. Por eso el agua recuperada nunca se elige por defecto: se verifica nivel, turbidez, conductividad y vigencia de la lectura."),
    ("4. Situación actual", "1:10", "La situación inicial tiene cuatro brechas: información dispersa, pérdida de excedentes, reúso sin criterio automatizado y baja trazabilidad. Estas brechas impiden calcular un balance y explicar una decisión después de un incidente."),
    ("5. Alternativa", "1:15", "La lectura entra autenticada, se valida y se compara con una regla. Si recuperada cumple nivel y calidad, se prioriza. Si falla una condición, se usa la principal y se genera una alerta. En ambas ramas se registra la causa."),
    ("6. Factibilidad", "1:00", "La alternativa es técnicamente factible con componentes disponibles; operacionalmente simple por perfiles; incremental en costo; y ambientalmente medible. La condición es calibrar sensores y validar hidráulica y agronomía antes del uso productivo."),
    ("7. Sistema", "1:20", "La arquitectura separa campo, API, datos y web. El ESP32 mantiene un watchdog local. Flask valida e impone roles. MongoDB conserva telemetría y trazabilidad. La web no decide por su cuenta: muestra la misma lógica del servidor."),
    ("8. Modelo", "1:05", "El modelo tiene diez colecciones. lecturas_sensor es time-series; usuarios y dispositivos contienen identidad protegida; reglas se versionan; eventos, alertas y auditoría reconstruyen qué pasó. Cada colección se exporta en JSON."),
    ("9. Requisitos funcionales", "1:00", "Los quince requisitos se agrupan en acceso, operación e información. Destaco RF08 y RF09, porque definen prioridad condicionada y estado seguro; RF10 registra control manual con motivo; y RF15 mantiene auditoría."),
    ("10. No funcionales", "1:10", "La seguridad incluye bcrypt, roles en servidor, CSRF, validación e idempotencia. El token de dispositivo se guarda como hash. HTTPS es una exigencia del despliegue productivo, no una característica que se simule en localhost."),
    ("11. Datos", "1:00", "La base demostrativa tiene 500 lecturas, catorce ciclos y diez de reúso. Los eventos suman 294 litros; 197 son reutilizados. El indicador calculado es 67,01 por ciento. Es un resultado del dataset, no una proyección anual."),
    ("12. Aplicación", "1:10", "En Administración se implementaron tres CRUD completos: usuarios, zonas de cultivo y dispositivos. Cada módulo permite registrar, consultar, editar y eliminar mediante baja lógica; se conserva el historial. El operador no tiene acceso a estas acciones y la validación se aplica en el servidor."),
    ("13. Pruebas", "1:30", "Se ejecutan dieciséis casos independientes. Cada caso vuelve a sembrar MongoDB, realiza la acción, verifica la respuesta y guarda un resultado. Las evidencias EV01 a EV16 se construyen desde esa ejecución y no desde resultados escritos previamente."),
    ("14. Conclusiones", "1:10", "El aporte es cerrar la cadena medir, validar, decidir, registrar y explicar. El reúso queda condicionado, la seguridad se aplica en capas y la entrega es reproducible. El paso siguiente es un piloto físico acotado con calibración y registro paralelo."),
    ("15. Cierre", "0:35", "En síntesis, AgroVic AquaCircular demuestra una forma responsable de convertir telemetría en decisiones hídricas. Ahora mostraré el flujo completo: ingreso, panel, perfiles, formularios, alerta, balance, base exportada y evidencias."),
]


def generate_defense_script(summary):
    document = Document()
    configure_docx(document)
    document.core_properties.author = "Sebastián Tapia Véliz"
    document.add_heading("Guion de defensa · AgroVic AquaCircular", level=0)
    document.add_paragraph("Sebastián Tapia Véliz · Proyecto de Título · septiembre de 2026")
    document.add_paragraph(
        "Duración estimada: 16–18 minutos de exposición y 10 minutos de demostración. El texto es una guía oral; conviene hablar mirando al jurado y usar la diapositiva como apoyo."
    )
    for title, duration, speech in SCRIPT:
        document.add_heading(f"{title} · {duration}", level=1)
        document.add_paragraph(speech)
    document.add_page_break()
    document.add_heading("Demostración funcional · 10 minutos", level=1)
    demo = [
        ("0:00–1:00", "Abrir repositorio público y mostrar README, carpetas y estado de Actions."),
        ("1:00–2:00", "Abrir CREDENCIALES_DEMO.txt local sin leer la contraseña en voz alta."),
        ("2:00–3:30", "Ingresar como Administrador y explicar las cuatro métricas del panel."),
        ("3:30–5:00", "Mostrar en Administración los tres CRUD: alta, consulta, modificar y desactivar registros de prueba, sin intervenir registros originales."),
        ("5:00–6:15", "Abrir otra sesión como Operador y comprobar que Administración responde 403."),
        ("6:15–7:30", "Enviar una lectura con simulador normal; después un escenario de calidad alta y observar la fuente principal/alerta."),
        ("7:30–8:30", "Registrar un riego manual con motivo y mostrar su auditoría."),
        ("8:30–9:15", "Abrir Balance y descargar el CSV con 294 L, 197 L reutilizados y 67,01%."),
        ("9:15–10:00", "Mostrar test_results.json, EV01–EV16 y los JSON exportables."),
    ]
    docx_table(document, ["Tiempo", "Acción y relato"], demo)
    document.add_heading("Preguntas probables", level=1)
    questions = [
        ("¿Está instalado en terreno?", "No. La evidencia cubre software, firmware de referencia y simulación reproducible. La instalación física requiere calibración y responsables de seguridad."),
        ("¿Por qué MongoDB?", "La telemetría es temporal y extensible. MongoDB permite meta por zona/dispositivo, colección time-series y exportación directa."),
        ("¿Qué ocurre sin internet?", "El firmware detiene la bomba localmente. El servidor marca desconexión y genera una alerta cuando vuelve a evaluar."),
        ("¿Cómo se evita reutilizar agua inadecuada?", "La regla exige nivel, calidad válida, turbidez y conductividad dentro de máximo. Si falla, se bloquea recuperada y se usa respaldo."),
        ("¿El 67,01% es ahorro real?", "No es una medición productiva; es el indicador exacto del conjunto de prueba para demostrar cálculo y trazabilidad."),
        ("¿Qué diferencia a los perfiles?", "Operador supervisa, confirma y ejecuta control autorizado; Administrador gestiona los CRUD de usuarios, zonas y equipos, reglas y cierres críticos."),
        ("¿Qué mejoraría primero?", "Pilotaje físico en una zona, calibración, validación agronómica y comparación con registro manual."),
    ]
    docx_table(document, ["Pregunta", "Respuesta breve"], questions)
    document.add_heading("Lista previa", level=1)
    for item in [
        "Aplicación y MongoDB levantados.",
        "Dos navegadores o perfiles listos.",
        "Credenciales locales verificadas.",
        "Simulador preparado.",
        "CSV, JSON y evidencias disponibles.",
        "Presentación en modo expositor y copia PDF del informe.",
        "No prometer instalación física ni resultados productivos.",
    ]:
        document.add_paragraph(item, style="List Bullet")
    docx_path = DIST / "Guion_Defensa_AgroVic_AquaCircular_Sebastian_Tapia.docx"
    document.save(docx_path)

    lines = ["GUION DE DEFENSA — AGROVIC AQUACIRCULAR", "Sebastián Tapia Véliz", ""]
    for title, duration, speech in SCRIPT:
        lines.extend([f"{title} ({duration})", speech, ""])
    lines.append("DEMOSTRACIÓN FUNCIONAL — 10 MINUTOS")
    for duration, action in demo:
        lines.append(f"{duration}: {action}")
    lines.extend(["", f"Resultado de pruebas: {summary['passed']}/{summary['total']} ({summary['pass_rate']}%)."])
    txt_path = DIST / "Guion_Defensa_AgroVic_AquaCircular_Sebastian_Tapia.txt"
    txt_path.write_text("\n".join(lines), encoding="utf-8")
    return docx_path, txt_path


def traceability_csv(summary):
    path = DIST / "Matriz_Trazabilidad_CP_EV.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Caso", "Evidencia", "Propósito", "Estado", "Requisitos", "Resultado observado"])
        for item in summary["results"]:
            writer.writerow([item["case"], item["evidence"], item["title"], item["status"], ";".join(item["requirements"]), item["observed"]])
    return path


def write_manifest(summary):
    path = DIST / "LEEME_PRIMERO.txt"
    path.write_text(
        "AGROVIC AQUACIRCULAR — ENTREGA FINAL\n"
        "Sebastián Tapia Véliz · Técnico en Informática · septiembre de 2026\n\n"
        "1. Informe final: Informe_Final_AgroVic_AquaCircular_Sebastian_Tapia.docx y .pdf\n"
        "2. Presentación: Presentacion_AgroVic_AquaCircular_Sebastian_Tapia.pptx (15 diapositivas)\n"
        "3. Guion: Guion_Defensa_AgroVic_AquaCircular_Sebastian_Tapia.docx y .txt\n"
        "4. Aplicación: backend/, frontend/, docker-compose.yml\n"
        "5. Base de datos: database/exports/ y database/README.md\n"
        "6. Evidencias: artifacts/evidence/ y artifacts/test_results.json\n"
        "7. Hardware/simulador: hardware/ y simulator/\n\n"
        f"Resultado automático: {summary['passed']}/{summary['total']} casos aprobados ({summary['pass_rate']}%).\n"
        "La instalación física en terreno no forma parte de la evidencia presentada.\n",
        encoding="utf-8",
    )
    return path


def checksums():
    path = DIST / "checksums_sha256.txt"
    rows = []
    for file in sorted(DIST.iterdir()):
        if file.is_file() and file.name not in {path.name} and file.suffix != ".zip":
            rows.append(f"{hashlib.sha256(file.read_bytes()).hexdigest()}  {file.name}")
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")
    return path


def package():
    target = DIST / "Entrega_Final_AgroVic_AquaCircular_Sebastian_Tapia.zip"
    include_dirs = ["backend", "frontend", "database", "hardware", "simulator", "tests", "docs", ".github"]
    include_root = ["README.md", "docker-compose.yml", ".gitignore"]
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        prefix = Path("AgroVic_AquaCircular")
        for name in include_root:
            path = ROOT / name
            if path.exists():
                archive.write(path, prefix / name)
        for directory in include_dirs:
            base = ROOT / directory
            if not base.exists():
                continue
            for path in base.rglob("*"):
                if path.is_file() and "__pycache__" not in path.parts:
                    archive.write(path, prefix / path.relative_to(ROOT))
        for path in DIST.iterdir():
            if path.is_file() and path != target:
                archive.write(path, prefix / "documentacion" / path.name)
        for path in [ARTIFACTS / "test_results.json", ARTIFACTS / "junit.xml", ARTIFACTS / "server.log"]:
            if path.exists():
                archive.write(path, prefix / path.relative_to(ROOT))
        if EVIDENCE.exists():
            for path in EVIDENCE.rglob("*"):
                if path.is_file():
                    archive.write(path, prefix / path.relative_to(ROOT))
        if DIAGRAMS.exists():
            for path in DIAGRAMS.rglob("*"):
                if path.is_file():
                    archive.write(path, prefix / path.relative_to(ROOT))
    return target


def main():
    DIST.mkdir(parents=True, exist_ok=True)
    summary = load_summary()
    diagrams = {
        "architecture": architecture_diagram(),
        "data": data_model_diagram(),
        "flow": flow_diagram(),
    }
    outputs = [
        generate_report(summary, diagrams),
        generate_report_pdf(summary, diagrams),
        generate_presentation(summary, diagrams),
        *generate_defense_script(summary),
        traceability_csv(summary),
        write_manifest(summary),
    ]
    outputs.append(checksums())
    archive = package()
    print(json.dumps({"generated": [str(path) for path in outputs], "archive": str(archive)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
